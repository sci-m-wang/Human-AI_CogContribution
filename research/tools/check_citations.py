import docx, re, sys, os
sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
d = docx.Document(os.path.join(ROOT, "paper", "ICIS_PDW_paper.docx"))
body, refs, seen = [], [], False
for p in d.paragraphs:
    t = p.text.strip()
    if not t: continue
    if p.style.name == "Heading 1" and t == "References": seen = True; continue
    (refs if seen else body).append(t)
text = " ".join(body)
# tables and captions are part of the body too
for tb in d.tables:
    for r in tb.rows:
        for c in r.cells: text += " " + c.text
# references: family name of first author (or organisation) and year
R = []
for r in refs:
    m = re.match(r"^(.+?)\.? \((\d{4})\)", r)
    first = re.split(r",", m.group(1))[0] if m else r[:20]
    R.append((first.strip(), m.group(2) if m else "?", r[:70]))
print("references:", len(R))
unc = []
for fam, yr, s in R:
    key = fam.split()[0] if fam.startswith("ICIS") else fam
    ok = re.search(re.escape(key) + r"[^()]{0,60}[(, ]\s*" + yr, text) or re.search(re.escape(key) + r".{0,120}" + yr, text)
    if not ok: unc.append((fam, yr, s))
print("references not cited in the body:", unc)
# citations in the body
cites = set()
for m in re.finditer(r"([A-Z][A-Za-z\u00C0-\u017F'’-]+)(?: et al\.| (?:&|and) [A-Z][A-Za-z\u00C0-\u017F'’-]+)?,? \(?((?:19|20)\d\d)[a-z]?", text):
    cites.add((m.group(1), m.group(2)))
famyr = {(f.split()[0] if f.startswith("ICIS") else f, y) for f, y, _ in R}
missing = sorted(c for c in cites if c not in famyr and c[0] not in ("Table", "Figure", "Version", "Appendix"))
print("citation-like strings with no matching reference:", missing)
