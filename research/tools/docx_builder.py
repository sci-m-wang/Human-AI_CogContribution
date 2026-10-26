import docx, re
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
HEADER_TITLE = 'Who Contributes What in Human-AI Research'
TEMPLATE = r'<project>\ICIS-2026_-Paper-Submission-Template-v3.docx'

def clear_body(d):
    body = d.element.body
    for el in list(body):
        if el.tag.endswith('sectPr'): continue
        body.remove(el)

def add_runs(p, text, size=None):
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for part in parts:
        if not part: continue
        if part.startswith('**'): r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('*'): r = p.add_run(part[1:-1]); r.italic = True
        else: r = p.add_run(part)
        if size: r.font.size = Pt(size)

def borders(table, color='7F7F7F', sz=4):
    """Kept for the legacy grid look. The papers now use three_line()."""
    tbl = table._tbl; pr = tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + e); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:space'), '0'); el.set(qn('w:color'), color); b.append(el)
    pr.append(b)

def three_line(table, rowlines=False, nheader=1):
    """Three-line table: heavy rule above and below, lighter rule under the header row, no vertical rules.
    rowlines=True adds hairlines between body rows (derivative form for tables with wrapped text)."""
    tbl = table._tbl; pr = tbl.tblPr
    b = OxmlElement('w:tblBorders')
    spec = [('top', 'single', 12, '000000'), ('left', 'nil', 0, 'auto'), ('bottom', 'single', 12, '000000'), ('right', 'nil', 0, 'auto'),
            ('insideH', 'single' if rowlines else 'nil', 2, 'A6A6A6'), ('insideV', 'nil', 0, 'auto')]
    for e, val, sz, col in spec:
        el = OxmlElement('w:' + e); el.set(qn('w:val'), val)
        if val != 'nil': el.set(qn('w:sz'), str(sz)); el.set(qn('w:space'), '0'); el.set(qn('w:color'), col)
        b.append(el)
    pr.append(b)
    for r in range(nheader):
        for c in table.rows[r].cells:
            tcPr = c._tc.get_or_add_tcPr(); tb = OxmlElement('w:tcBorders')
            bt = OxmlElement('w:bottom'); bt.set(qn('w:val'), 'single'); bt.set(qn('w:sz'), '6'); bt.set(qn('w:space'), '0'); bt.set(qn('w:color'), '000000'); tb.append(bt)
            tcPr.append(tb)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); tcPr.append(s)

def set_cell(cell, text, bold=False, size=8.5, align='left'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1); p.paragraph_format.line_spacing = 1.0
    p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER}[align]
    add_runs(p, text, size)
    if bold:
        for r in p.runs: r.bold = True

def box(d, lines, fill='EAF1F8', edge='44709E', size=10):
    """Set-off block: one shaded cell with a heavy left rule. lines = list of (text, indent_inches)."""
    t = d.add_table(rows=1, cols=1); t.autofit = False
    tbl = t._tbl; pr = tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e, sz, col in (('top', 4, edge), ('left', 24, edge), ('bottom', 4, edge), ('right', 4, edge)):
        el = OxmlElement('w:' + e); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:space'), '0'); el.set(qn('w:color'), col); b.append(el)
    pr.append(b)
    cell = t.cell(0, 0); cell.width = Inches(6.5); shade(cell, fill)
    mar = OxmlElement('w:tcMar')
    for side, w in (('top', 80), ('left', 160), ('bottom', 80), ('right', 120)):
        m = OxmlElement('w:' + side); m.set(qn('w:w'), str(w)); m.set(qn('w:type'), 'dxa'); mar.append(m)
    cell._tc.get_or_add_tcPr().append(mar)
    for i, (text, indent) in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(3); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.0
        if indent:
            p.paragraph_format.left_indent = Inches(indent); p.paragraph_format.first_line_indent = Inches(-indent)
        add_runs(p, text, size)
    sp = d.add_paragraph(style='Normal'); sp.paragraph_format.space_after = Pt(0); sp.paragraph_format.space_before = Pt(0)
    for r in sp.runs: r.font.size = Pt(3)
    sp.paragraph_format.line_spacing = Pt(5)

def build(blocks, out):
    d = docx.Document(TEMPLATE)
    clear_body(d)
    for sec in d.sections:
        for hp in sec.header.paragraphs:
            if 'Short Title' in hp.text:
                for r in hp.runs: r.text=''
                hp.runs[0].text=HEADER_TITLE
    for b in blocks:
        k = b[0]
        if k == 'title': d.add_paragraph(b[1], style='Title')
        elif k == 'h1': d.add_paragraph(b[1], style='Heading 1')
        elif k == 'h2': d.add_paragraph(b[1], style='Heading 2')
        elif k == 'p':
            p = d.add_paragraph(style='Normal'); add_runs(p, b[1]); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif k == 'ref':
            p = d.add_paragraph(style='Normal'); add_runs(p, b[1]); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.3); p.paragraph_format.space_after = Pt(2)
        elif k == 'caption':
            p = d.add_paragraph(style='Caption'); add_runs(p, b[1]); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = False
        elif k == 'table':
            rows = b[1]; widths = b[2] if len(b) > 2 else None
            t = d.add_table(rows=len(rows), cols=len(rows[0])); t.autofit = False
            three_line(t, rowlines=(max(len(str(v)) for r_ in rows for v in r_) > 70))
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    set_cell(t.cell(i, j), str(val), bold=(i == 0), align=('left' if j == 0 or (widths and widths[j] > 1.6) else 'center'))
                    for pp in t.cell(i, j).paragraphs: pp.paragraph_format.keep_with_next = (len(rows) <= 8) or i <= 1 or i >= len(rows) - 2
                    if widths: t.cell(i, j).width = Inches(widths[j])
                _trPr = t.rows[i]._tr.get_or_add_trPr()
                _trPr.append(OxmlElement('w:cantSplit'))
                if i == 0: _trPr.append(OxmlElement('w:tblHeader'))
        elif k == 'image':
            d.add_paragraph(style='Normal').alignment = WD_ALIGN_PARAGRAPH.CENTER
            d.paragraphs[-1].add_run().add_picture(b[1], width=Inches(b[2])); d.paragraphs[-1].paragraph_format.keep_with_next = True
        elif k == 'bullets':
            for n_, item in enumerate(b[1]):
                p = d.add_paragraph(style='Normal'); add_runs(p, '\u2022\t' + item); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                pf = p.paragraph_format; pf.left_indent = Inches(0.3); pf.first_line_indent = Inches(-0.2); pf.tab_stops.add_tab_stop(Inches(0.3))
                pf.space_before = Pt(0); pf.space_after = Pt(6 if n_ == len(b[1]) - 1 else 2)
        elif k == 'box':
            box(d, b[1])
        elif k == 'h1pb':
            hp = d.add_paragraph(b[1], style='Heading 1'); hp.paragraph_format.page_break_before = True
        elif k == 'pagebreak':
            d.add_page_break()
    d.save(out)
