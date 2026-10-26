import docx, sys
d = docx.Document(r'<project>\ICIS-2026_-Paper-Submission-Template-v3.docx')
sec = d.sections[0]
print('page', sec.page_width.inches, sec.page_height.inches, 'margins', sec.left_margin.inches, sec.right_margin.inches, sec.top_margin.inches, sec.bottom_margin.inches)
print('styles used:')
used = {}
for p in d.paragraphs:
    used.setdefault(p.style.name, 0); used[p.style.name]+=1
print(used)
for i,p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t: print(i, '[%s]'%p.style.name, t[:160])
print('tables', len(d.tables))
